import os
import win32com.client
import pythoncom
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

def create_massive_report():
    save_dir = r"C:\Users\BELIKANON\Desktop"
    docx_path = os.path.join(save_dir, "Test_Rapport_Audit_Kibali.docx")
    pdf_path = os.path.join(save_dir, "Test_Rapport_Audit_Kibali.pdf")

    print("📝 1. Génération du document Word (15 pages)...")
    doc = Document()
    
    # Page de garde
    title = doc.add_heading('RAPPORT D\'AUDIT FINANCIER ET FISCAL', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph('\n\nVAALCO GABON (2005 - 2006)')
    doc.add_page_break()

    # Contenu sur 15 pages (simulé avec 15 sections)
    for i in range(1, 16):
        doc.add_heading(f'Section {i} : Analyse de la période', level=1)
        
        paragraph = doc.add_paragraph(
            f'Dans le cadre de notre mission d\'audit pour l\'exercice {2005 if i%2==0 else 2006}, '
            'nous avons procédé à l\'examen des états financiers de la société VAALCO GABON. '
            'Les calculs de taxes et d\'impôts ont été réalisés sur la base du Code Général des Impôts en vigueur. '
            'Les résultats présentent une conformité satisfaisante, sous réserve des ajustements mentionnés.\n\n'
        )
        paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        paragraph.style.font.size = Pt(11)
        
        # Ajout d'un faux tableau de chiffres pour faire pro
        table = doc.add_table(rows=3, cols=3)
        table.style = 'Table Grid'
        hdr_cells = table.rows[0].cells
        hdr_cells[0].text = 'Libellé'
        hdr_cells[1].text = 'Montant HT'
        hdr_cells[2].text = 'TVA (18%)'
        row_cells = table.rows[1].cells
        row_cells[0].text = 'Chiffre d\'Affaires'
        row_cells[1].text = '4 120 684 USD'
        row_cells[2].text = '741 723 USD'
        
        doc.add_page_break()

    doc.save(docx_path)
    print(f"✅ Word généré : {docx_path}")

    print("📄 2. Conversion en PDF via Word COM...")
    pythoncom.CoInitialize()
    try:
        word = win32com.client.Dispatch("Word.Application")
        word.Visible = False
        doc_com = word.Documents.Open(docx_path)
        
        # 17 = wdFormatPDF
        doc_com.SaveAs2(pdf_path, FileFormat=17)
        doc_com.Close(False)
        word.Quit()
        print(f"✅ PDF généré : {pdf_path}")
    except Exception as e:
        print(f"❌ Erreur COM : {e}")

if __name__ == "__main__":
    create_massive_report()