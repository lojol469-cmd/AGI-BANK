CURRENT_DATE = "2026-10-28"

import json
import os
import sys

clean_input = "{\"filename\": \"Rapport_Kibali\", \"section_title\": \"Bilan 2026\", \"content\": \"Tout fonctionne parfaitement\"}".replace("```json", "").replace("```", "").strip()
try:
    payload = json.loads(clean_input)
except json.JSONDecodeError:
    print(json.dumps({"success": False, "output": "❌ JSON invalide."}))
    sys.exit(0)

try:
    from docx import Document
    from docx.shared import Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT
except ImportError:
    print(json.dumps({"success": False, "output": "❌ python-docx manquant. (Utilise ultra_lib_installer)"}))
    sys.exit(0)

filename = payload.get("filename", "Rapport_Audit")
save_dir = os.path.join(os.path.expanduser("~"), "Desktop")
docx_path = os.path.join(save_dir, f"{filename}.docx")

# CAS 1 : FINALISATION
if payload.get("finalize"):
    pdf_path = os.path.join(save_dir, f"{filename}.pdf")
    try:
        import win32com.client
        import pythoncom
    except ImportError:
        print(json.dumps({"success": False, "output": "❌ pywin32 manquant pour le PDF. Le Word a été créé."}))
        sys.exit(0)
        
    try:
        pythoncom.CoInitialize()
        word = win32com.client.Dispatch("Word.Application")
        word.Visible = False
        doc = word.Documents.Open(docx_path)
        doc.SaveAs2(pdf_path, FileFormat=17)
        doc.Close(False)
        word.Quit()
        print(json.dumps({"success": True, "output": f"✅ Rapport finalisé !\nWord: {docx_path}\nPDF: {pdf_path}"}))
    except Exception as e:
        print(json.dumps({"success": False, "output": f"❌ Erreur COM : {e}"}))
    sys.exit(0)

# CAS 2 : AJOUT D'UNE SECTION
section_title = payload.get("section_title", "")
content = payload.get("content", "")
table_data = payload.get("table", None)

try:
    if os.path.exists(docx_path):
        doc = Document(docx_path)
    else:
        doc = Document()
        title = doc.add_heading("RAPPORT D'AUDIT FINANCIER ET FISCAL", 0)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.add_paragraph("\n\nDossier VAALCO GABON")
        doc.add_page_break()

    if section_title:
        doc.add_heading(section_title, level=1)

    if content:
        p = doc.add_paragraph(content)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.style.font.size = Pt(11)

    if table_data and isinstance(table_data, list) and len(table_data) > 0:
        rows = len(table_data)
        cols = len(table_data[0])
        table = doc.add_table(rows=rows, cols=cols)
        table.style = 'Table Grid'
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        
        for i, row in enumerate(table_data):
            for j, text in enumerate(row):
                cell = table.cell(i, j)
                cell.text = str(text)
                if i == 0:
                    cell.paragraphs[0].runs[0].bold = True
        doc.add_paragraph()

    doc.add_page_break()
    doc.save(docx_path)
    print(json.dumps({"success": True, "output": f"✅ Section '{section_title}' ajoutée avec succès."}))
except Exception as e:
    print(json.dumps({"success": False, "output": f"❌ Erreur d'écriture : {str(e)}"}))
